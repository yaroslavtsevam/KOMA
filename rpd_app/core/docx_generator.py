#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx_generator.py
Универсальный генератор итоговой рабочей программы (РПД)
с пост-процессингом таблиц, стилей, слияний ячеек и валидацией.
"""

import os
import re
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from docxtpl import DocxTemplate


def remove_element(elem):
    parent = elem.getparent()
    if parent is not None:
        parent.remove(elem)


def postprocess_table_1(doc):
    """Таблица 1: Вертикальное объединение ячеек компетенций через w:vMerge."""
    t1 = doc.tables[1]
    w_ns = nsdecls("w")
    current_code = None
    restart_row_idx = None

    for r_idx in range(2, len(t1.rows)):
        row = t1.rows[r_idx]
        c0_text = row.cells[0].text.strip()
        c1_text = row.cells[1].text.strip()
        
        # Определяем код компетенции из индикатора или кода
        comp_code = None
        for txt in [c1_text, c0_text]:
            m = re.search(r'^(УК|ОПК|ПК|ПКос)-[\d]+', txt)
            if m:
                comp_code = m.group(0)
                break

        if comp_code != current_code:
            current_code = comp_code
            restart_row_idx = r_idx
            for c_idx in range(4):
                tc = row.cells[c_idx]._tc
                tcPr = tc.get_or_add_tcPr()
                for old in tcPr.xpath('./w:vMerge'):
                    tcPr.remove(old)
                tcPr.append(parse_xml(f'<w:vMerge {w_ns} w:val="restart"/>'))
        else:
            for c_idx in range(4):
                tc = row.cells[c_idx]._tc
                tcPr = tc.get_or_add_tcPr()
                for old in tcPr.xpath('./w:vMerge'):
                    tcPr.remove(old)
                tcPr.append(parse_xml(f'<w:vMerge {w_ns}/>'))


def sanitize_section2_text(context: dict) -> dict:
    """Очищает поля Раздела 2 от задвоения вводных фраз шаблона."""
    if not isinstance(context, dict):
        return context

    # 1. prerequisites_text
    prereq = context.get("prerequisites_text", "")
    if isinstance(prereq, str) and prereq:
        prereq = re.sub(
            r"^(?:Освоение\s+дисциплины\s+(?:непосредственно\s+)?базируется\s+на\s+(?:результатах\s+обучения\s+по\s+)?(?:предшествующим\s+дисциплинам|предшествующих\s+дисциплинах|курсах)?[:\s]*|"
            r"Предшествующими\s+(?:курсами|дисциплинами)(?:,\s*на\s+которых\s+базируется\s+дисциплина,)?\s+являются[:\s]*|"
            r"К\s+предшествующим\s+(?:курсам|дисциплинам)\s+относятся[:\s]*|"
            r"В\s+качестве\s+пререквизитов\s+(?:выступают|используются)[:\s]*|"
            r"Изучение\s+дисциплины\s+базируется\s+на[:\s]*|"
            r"Базируется\s+на[:\s]*)",
            "",
            prereq,
            flags=re.IGNORECASE
        ).strip()
        context["prerequisites_text"] = prereq

    # 2. postrequisites_text
    postreq = context.get("postrequisites_text", "")
    if isinstance(postreq, str) and postreq:
        postreq = re.sub(
            r"^(?:Освоение\s+дисциплины\s+необходимо\s+для\s+(?:последующего\s+изучения\s+дисциплин\s+и\s+прохождения\s+практик|последующих\s+дисциплин)?[:\s]*|"
            r"Последующими\s+(?:курсами|дисциплинами)\s+являются[:\s]*|"
            r"К\s+последующим\s+(?:курсам|дисциплинам)\s+относятся[:\s]*|"
            r"Необходимо\s+для\s+(?:последующего\s+изучения|изучения)?[:\s]*|"
            r"В\s+качестве\s+постреквизитов\s+(?:выступают|используются)[:\s]*)",
            "",
            postreq,
            flags=re.IGNORECASE
        ).strip()
        context["postrequisites_text"] = postreq

    # 3. corequisites_text
    coreq = context.get("corequisites_text", "")
    if isinstance(coreq, str) and coreq:
        coreq = re.sub(
            r"^(?:(?:В\s+[^\n,:]+семестре\s+)?параллельно\s+изучаются\s+дисциплины[:\s]*|"
            r"Параллельно\s+изучаемыми\s+дисциплинами\s+являются[:\s]*|"
            r"К\s+параллельно\s+изучаемым\s+дисциплинам\s+относятся[:\s]*|"
            r"Дисциплины\s+параллельного\s+освоения[:\s]*)",
            "",
            coreq,
            flags=re.IGNORECASE
        ).strip()
        context["corequisites_text"] = coreq

    # 4. course_features_text
    feat = context.get("course_features_text", "")
    if isinstance(feat, str) and feat:
        feat = re.sub(
            r"^(?:Особенностью\s+дисциплины\s+является[:\s]*|"
            r"Особенности\s+дисциплины\s+заключаются\s+в\s+том,\s*что[:\s]*|"
            r"К\s+особенностям\s+дисциплины\s+относятся?[:\s]*|"
            r"Особенность\s+курса[:\s]*)",
            "",
            feat,
            flags=re.IGNORECASE
        ).strip()
        if re.match(r"^дисциплина\s+носит\s+", feat, flags=re.IGNORECASE):
            feat = re.sub(r"^дисциплина\s+носит\s+", "", feat, flags=re.IGNORECASE)
        if feat and feat[0].isupper() and not feat.startswith("«"):
            feat = feat[0].lower() + feat[1:]
        context["course_features_text"] = feat

    # 5. BRS assessment text
    use_brs = context.get("use_brs", True)
    if "brs_assessment_text" not in context or not context["brs_assessment_text"]:
        context["brs_assessment_text"] = (
            "Для аттестации обучающихся используется балльно-рейтинговая система (БРС) оценивания."
            if use_brs else
            "Для аттестации обучающихся используется традиционная система оценивания."
        )

    return context


def postprocess_table_2(doc, context=None):
    """Таблица 2: Удаление строк с нулевыми часами и устранение пустой 4-й колонки для 1-семестровых курсов."""
    t2 = doc.tables[2]
    rows_to_delete = []
    for r_idx in range(len(t2.rows) - 1, -1, -1):
        row = t2.rows[r_idx]
        name_cell = row.cells[0].text.strip()
        if any(h in name_cell for h in [
            "Вид учебной работы", "Общая трудоёмкость", "1. Контактная работа",
            "Аудиторная работа", "2. Самостоятельная работа", "Форма промежуточной",
            "Вид промежуточного контроля"
        ]):
            continue
        hour_vals = [c.text.strip().replace(' ', '') for c in row.cells[1:]]
        is_empty = all(v in ["", "0", "0,0", "0.0", "-", "–"] for v in hour_vals)
        if is_empty:
            rows_to_delete.append(r_idx)

    for r_idx in rows_to_delete:
        remove_element(t2.rows[r_idx]._tr)

    # Проверка, является ли курс односеместровым
    has_sem_2 = False
    if context and context.get("sem_2_hdr"):
        has_sem_2 = bool(str(context.get("sem_2_hdr")).strip())
    else:
        if len(t2.rows) > 1:
            tcs_r1 = t2.rows[1]._tr.xpath('w:tc')
            if len(tcs_r1) >= 4:
                txt_sem2 = "".join(tcs_r1[3].itertext()).strip()
                has_sem_2 = bool(txt_sem2)

    if not has_sem_2:
        # Односеместровый курс: удаляем 4-ю колонку и делаем чистую таблицу из 3 колонок
        grid = t2._tbl.tblGrid
        cols = grid.xpath('w:gridCol')
        if len(cols) == 4:
            grid.remove(cols[3])
            cols[0].set(qn('w:w'), '6800')
            cols[1].set(qn('w:w'), '2000')
            cols[2].set(qn('w:w'), '2000')

        # Row 0: ячейка 2 ("В т.ч. семестр") имела gridSpan=2, снимаем gridSpan
        if len(t2.rows) > 0:
            r0_tcs = t2.rows[0]._tr.xpath('w:tc')
            if len(r0_tcs) >= 3:
                gs0 = r0_tcs[2].xpath('.//w:gridSpan')
                if gs0:
                    gs0[0].getparent().remove(gs0[0])

        # Строка "в том числе:": уменьшаем gridSpan с 4 до 3
        for r in t2.rows:
            r_tcs = r._tr.xpath('w:tc')
            if len(r_tcs) == 1:
                gs = r_tcs[0].xpath('.//w:gridSpan')
                if gs:
                    gs[0].set(qn('w:val'), '3')

        # Во всех строках с 4 ячейками удаляем 4-ю ячейку
        for r in t2.rows:
            r_tcs = r._tr.xpath('w:tc')
            if len(r_tcs) == 4:
                r._tr.remove(r_tcs[3])
    else:
        # Для двухсеместрового курса: строка "Вид промежуточного контроля" должна иметь gridSpan=3
        for r in t2.rows:
            r_tcs = r._tr.xpath('w:tc')
            if len(r_tcs) == 2:
                txt_0 = "".join(r_tcs[0].itertext()).strip()
                if "промежуточного контроля" in txt_0:
                    gs = r_tcs[1].xpath('.//w:gridSpan')
                    if gs:
                        gs[0].set(qn('w:val'), '3')


def postprocess_table_3(doc):
    """Таблица 3: Сброс bold в строках данных (строки >= 2) и замена 0 на пустоту."""
    t3 = doc.tables[3]
    for row in t3.rows[2:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip() in ["0", "0,0", "0.0"]:
                    p.text = ""
                for r in p.runs:
                    r.bold = False


def postprocess_table_5(doc):
    """Таблица 5: Горизонтальное объединение ячеек в строках-разделителях разделов."""
    t5 = doc.tables[5]
    for row in t5.rows[1:]:
        c0 = row.cells[0].text.strip()
        c1 = row.cells[1].text.strip()
        c2 = row.cells[2].text.strip()
        if c0 and (c0 == c1 == c2) and c0.startswith("Раздел"):
            row.cells[0].merge(row.cells[1]).merge(row.cells[2])
            p = row.cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(11.0)


def postprocess_toc(doc):
    """Оглавление: Снятие синего цвета и подчеркивания у гиперссылок."""
    w_ns = nsdecls("w")
    for p in doc.paragraphs[50:90]:
        for h in p._p.xpath('.//w:hyperlink'):
            for r_node in h.xpath('.//w:r'):
                rPr = r_node.find(qn('w:rPr'))
                if rPr is not None:
                    col = rPr.find(qn('w:color'))
                    if col is not None:
                        rPr.remove(col)
                    u = rPr.find(qn('w:u'))
                    if u is not None:
                        rPr.remove(u)
                    rPr.append(parse_xml(f'<w:color {w_ns} w:val="auto"/>'))
                    rPr.append(parse_xml(f'<w:u {w_ns} w:val="none"/>'))


def fix_developer_block(doc):
    """Нормализация блока разработчиков: удаление (подпись), выравнивание даты, 14 кегль."""
    in_dev = False
    for p in doc.paragraphs[20:60]:
        t = p.text.strip()
        if t.startswith("Разработчик:") or t.startswith("Разработчики:"):
            in_dev = True
        if in_dev:
            if "(подпись)" in p.text:
                p.text = p.text.replace("(подпись)", "").strip()
            if "«__» ________" in p.text:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14.0)
            if "Программа обсуждена на заседании" in t:
                break


def validate_rpd_document(doc_path: str) -> dict:
    """Валидация сформированного документа на отсутствие нераскрытых тегов и базовую корректность."""
    doc = docx.Document(doc_path)
    unrendered = []
    for i, p in enumerate(doc.paragraphs):
        if "{{" in p.text or "{%" in p.text:
            unrendered.append(f"Параграф {i}: {p.text[:50]}")
    for ti, t in enumerate(doc.tables):
        for ri, r in enumerate(t.rows):
            for ci, c in enumerate(r.cells):
                if "{{" in c.text or "{%" in c.text:
                    unrendered.append(f"Таблица {ti} r{ri} c{ci}: {c.text[:40]}")

    landscape_ok = False
    if len(doc.sections) > 1:
        s1 = doc.sections[1]
        landscape_ok = (s1.orientation.name == "LANDSCAPE" or s1.page_width > s1.page_height)

    return {
        "valid": len(unrendered) == 0 and landscape_ok,
        "unrendered_tags": unrendered,
        "landscape_table_1": landscape_ok,
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables)
    }


def generate_rpd(template_path: str, context: dict, out_path: str) -> str:
    """
    Полный цикл генерации: рендеринг docxtpl -> пост-процессинг -> валидация -> сохранение.
    """
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)

    # Санитизация полей Раздела 2
    context = sanitize_section2_text(context)

    # 1. Рендеринг шаблона через docxtpl
    tpl = DocxTemplate(template_path)
    tpl.render(context)
    tpl.save(out_path)

    # 2. Пост-процессинг верстки в python-docx
    doc = docx.Document(out_path)
    postprocess_table_1(doc)
    postprocess_table_2(doc, context)
    postprocess_table_3(doc)
    postprocess_table_5(doc)
    postprocess_toc(doc)
    fix_developer_block(doc)

    doc.save(out_path)

    # 3. Валидация
    val = validate_rpd_document(out_path)
    if not val["valid"]:
        print("ВНИМАНИЕ: Обнаружены замечания валидации РПД:", val)

    return out_path
