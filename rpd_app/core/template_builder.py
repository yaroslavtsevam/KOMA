#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
template_builder.py
Универсальный строитель чистого параметризованного шаблона РПД
на основе проверенных OpenXML-преобразований.
"""

import os
import re
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls


def remove_element(elem):
    parent = elem.getparent()
    if parent is not None:
        parent.remove(elem)


def set_run_font(run, name="Times New Roman", size_pt=14.0, bold=None, italic=None, color=None):
    run.font.name = name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.xpath('./w:rFonts')
    if not rFonts:
        rf = OxmlElement('w:rFonts')
        rf.set(qn('w:ascii'), name)
        rf.set(qn('w:hAnsi'), name)
        rf.set(qn('w:cs'), name)
        rf.set(qn('w:eastAsia'), name)
        rPr.append(rf)
    else:
        rf = rFonts[0]
        rf.set(qn('w:ascii'), name)
        rf.set(qn('w:hAnsi'), name)
        rf.set(qn('w:cs'), name)


def set_para_format(p, alignment=None, line_spacing=1.0, space_before_pt=0, space_after_pt=0, first_line_indent_pt=None):
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    if first_line_indent_pt is not None:
        p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)


def create_paragraph_after(target_elem, doc, text="", style='Normal'):
    p_new = doc.add_paragraph(style=style)
    p_new.text = text
    target_elem.addnext(p_new._p)
    return p_new


def build_rpd_template(src_docx_path: str, out_docx_path: str):
    """
    Применяет полный комплекс проверенных правил верстки и нормализации OpenXML
    к исходному шаблону и сохраняет готовый rpd_template_parametrized.docx.
    """
    if not os.path.exists(src_docx_path):
        raise FileNotFoundError(f"Исходный шаблон не найден: {src_docx_path}")

    os.makedirs(os.path.dirname(os.path.abspath(out_docx_path)), exist_ok=True)
    doc = docx.Document(src_docx_path)
    w_ns = nsdecls("w")

    # 1. ТИТУЛЬНЫЙ ЛИСТ
    # Шапка ВУЗа (Таблица 0, ячейка 1)
    t0 = doc.tables[0]
    c_hdr = t0.rows[0].cells[1]
    while len(c_hdr.paragraphs) > 1:
        remove_element(c_hdr.paragraphs[-1]._p)
    p_h1 = c_hdr.paragraphs[0]
    p_h1.text = "МИНИСТЕРСТВО СЕЛЬСКОГО ХОЗЯЙСТВА РОССИЙСКОЙ ФЕДЕРАЦИИ"
    set_para_format(p_h1, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=1)
    for r in p_h1.runs:
        set_run_font(r, name="Times New Roman", size_pt=10.0, bold=True)

    # Строка 2: всеми заглавными буквами, 7 pt
    p_h2 = create_paragraph_after(p_h1._p, doc, text="ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ")
    set_para_format(p_h2, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=1)
    for r in p_h2.runs:
        set_run_font(r, name="Times New Roman", size_pt=7.0, bold=False)

    p_h3 = create_paragraph_after(p_h2._p, doc, text="«РОССИЙСКИЙ ГОСУДАРСТВЕННЫЙ АГРАРНЫЙ УНИВЕРСИТЕТ – МСХА имени К.А. ТИМИРЯЗЕВА»")
    set_para_format(p_h3, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=1)
    for r in p_h3.runs:
        set_run_font(r, name="Times New Roman", size_pt=10.0, bold=True)

    p_h4 = create_paragraph_after(p_h3._p, doc, text="(ФГБОУ ВО РГАУ - МСХА имени К.А. Тимирязева)")
    set_para_format(p_h4, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=0)
    for r in p_h4.runs:
        set_run_font(r, name="Times New Roman", size_pt=10.0, bold=True)

    # Параметризация курса, семестра и отступов
    for p in doc.paragraphs[:30]:
        t = p.text.strip()
        if "РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ" in t:
            p.text = "РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ"
            set_para_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before_pt=12, space_after_pt=12)
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size_pt=14.0, bold=True)
        elif t.startswith("Курс ") or t.startswith("Курс {{"):
            p.text = "Курс {{ course_year }}"
            set_para_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_pt=0)
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size_pt=14.0)
        elif t.startswith("Семестр ") or t.startswith("Семестр {{"):
            p.text = "Семестр {{ semester }}"
            set_para_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_pt=0)
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size_pt=14.0)

    # 2. ВТОРАЯ СТРАНИЦА (РАЗРАБОТЧИКИ И КАФЕДРА)
    for p in doc.paragraphs[20:60]:
        t = p.text.strip()
        if "Программа обсуждена на заседании" in t:
            p.text = "Программа обсуждена на заседании кафедры {{ department_name_short }}, протокол от «__» ________ {{ current_year }} г. № ___"
            set_para_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_pt=0)
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size_pt=14.0)
        elif "выпускающей кафедрой" in t:
            p.text = "Согласовано с выпускающей кафедрой {{ graduating_department_name_short }}, протокол от «__» ________ {{ current_year }} г. № ___"
            set_para_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_pt=0)
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size_pt=14.0)

    # 3. ОГЛАВЛЕНИЕ (TOC) — Очистка от синего цвета и подчеркиваний
    for p in doc.paragraphs[50:90]:
        for h in p._p.xpath('.//w:hyperlink'):
            for r_node in h.xpath('.//w:r'):
                rPr = r_node.find(qn('w:rPr'))
                if rPr is not None:
                    col = rPr.find(qn('w:color'))
                    if col is not None:
                        rPr.remove(col)
                    u = rPr.find(qn('w:u'))
                    if u is not None:
                        rPr.remove(u)
                    rPr.append(parse_xml(f'<w:color {w_ns} w:val="auto"/>'))
                    rPr.append(parse_xml(f'<w:u {w_ns} w:val="none"/>'))

    # 4. ИЗОЛЯЦИЯ ТАБЛИЦЫ 1 В АЛЬБОМНЫЙ РАЗДЕЛ (LANDSCAPE)
    body = doc._body._body
    el_sec3_intro = None
    el_sec4 = None
    el_sec41 = None
    el_workload = None
    el_t1_label = None

    for i, el in enumerate(body):
        if i < 70:
            continue
        txt = "".join(el.itertext()).strip()
        tag = el.tag.split("}")[-1]
        if tag == "p":
            if "Изучение данной учебной дисциплины направлено" in txt or "3. Перечень планируемых результатов" in txt:
                if "Изучение данной" in txt:
                    el_sec3_intro = el
            elif (txt.startswith("4. Структура") or txt.startswith("4.  Структура")) and el_sec4 is None:
                el_sec4 = el
            elif txt.startswith("4.1") and el_sec41 is None:
                el_sec41 = el
            elif txt.startswith("Общая трудоёмкость дисциплины составляет") and el_workload is None:
                el_workload = el
            elif (txt.startswith("Таблица 1") and not "Таблица 10" in txt) and el_t1_label is None:
                el_t1_label = el

    if el_t1_label is not None:
        el_t1_title = el_t1_label.getnext()
        el_t1_tbl = el_t1_title.getnext()

    sectPrs = body.xpath(".//w:sectPr")
    el_p_port = sectPrs[0].getparent().getparent() if len(sectPrs) > 0 else None
    el_p_land = sectPrs[1].getparent().getparent() if len(sectPrs) > 1 else None

    if all(x is not None for x in [el_sec3_intro, el_sec4, el_sec41, el_workload, el_t1_label, el_t1_title, el_t1_tbl, el_p_port, el_p_land]):
        p_blank_t1 = parse_xml(f"<w:p {w_ns}/>")
        el_sec3_intro.addnext(el_p_port)
        el_p_port.addnext(p_blank_t1)
        p_blank_t1.addnext(el_t1_label)
        el_t1_label.addnext(el_t1_title)
        el_t1_title.addnext(el_t1_tbl)
        el_t1_tbl.addnext(el_p_land)
        el_p_land.addnext(el_sec4)
        el_sec4.addnext(el_workload)
        el_workload.addnext(el_sec41)

    # 5. ТАБЛИЦА 3: ТОЛЬКО ШАПКА ЖИРНАЯ (УДАЛЕНИЕ BOLD ИЗ ДАННЫХ)
    t3 = doc.tables[3]
    for row in t3.rows[2:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = False

    # 6. ВСТАВКА ПУСТОЙ СТРОКИ ПЕРЕД ВСЕМИ ТАБЛИЦАМИ 1-10
    tbl_pattern = re.compile(r"^Таблица\s+(\d+)", re.IGNORECASE)
    table_captions = {}
    for el in doc._body._body:
        if el.tag.endswith("}p"):
            txt = "".join(el.itertext()).strip()
            m = tbl_pattern.match(txt)
            if m:
                num = int(m.group(1))
                if num not in table_captions:
                    table_captions[num] = el

    for num in range(1, 11):
        if num in table_captions:
            cap_el = table_captions[num]
            prev = cap_el.getprevious()
            needs_blank = True
            if prev is not None and prev.tag.endswith("}p"):
                pPr = prev.find(qn("w:pPr"))
                has_sectPr = pPr is not None and pPr.find(qn("w:sectPr")) is not None
                txt = "".join(prev.itertext()).strip()
                if not has_sectPr and txt == "":
                    needs_blank = False
            if needs_blank:
                p_blank = parse_xml(f"<w:p {w_ns}/>")
                cap_el.addprevious(p_blank)

    doc.save(out_docx_path)
    return out_docx_path
